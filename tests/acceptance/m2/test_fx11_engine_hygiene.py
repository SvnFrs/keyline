"""Audit 03 FX-11 (amendment B-19): engine output that is not UTF-8, a validate timeout,
SIGINT and SIGTERM mid-render and mid-validate, and an empty -o.

Each signal test runs keyline in a subprocess with a private TMPDIR, sends the signal
once an engine child is running, and then checks that keyline exited 130 or 143, that
no keyline temp directory is left, and that no engine child survives. Shell stand-ins
for soffice and officecli make the timing deterministic and run anywhere; the real
engines are tested where they are installed."""

import json
import os
import shutil
import signal
import subprocess
import sys
import time
from pathlib import Path

import pytest

from keyline import render as render_mod
from keyline import validate as validate_mod
from tests.acceptance._cli import keyline
from tests.conftest import GOLDEN

pytestmark = pytest.mark.skipif(sys.platform != "linux", reason="process groups and /proc")

CLEAN = GOLDEN / "editorial-fixed.pptx"
KPI = GOLDEN / "kpi-recipe.pptx"
SLEEP = shutil.which("sleep")
# the stand-in engines start a real `sleep`; with no PATH (the empty-PATH run) they cannot
needs_sleep = pytest.mark.skipif(SLEEP is None, reason="no sleep binary for the stand-ins")


def _stand_in(bin_dir: Path, name: str, pidfile: Path, version: str) -> None:
    """An engine that answers --version and close, and otherwise sleeps in a child."""
    bin_dir.mkdir(exist_ok=True)
    exe = bin_dir / name
    exe.write_text(
        "#!/bin/sh\n"
        f'case "$1" in --version) echo "{version}"; exit 0;; close) exit 0;; esac\n'
        f"{SLEEP} 300 &\n"
        f"echo $! > {pidfile}\n"
        "wait\n"
    )
    exe.chmod(0o755)


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    status = Path(f"/proc/{pid}/status")
    return status.exists() and "State:\tZ" not in status.read_text()


def _mentioning(text: str) -> list[int]:
    """Live processes whose command line contains `text`."""
    found = []
    for d in Path("/proc").iterdir():
        if d.name.isdigit() and int(d.name) != os.getpid():
            try:
                if text.encode() in (d / "cmdline").read_bytes():
                    found.append(int(d.name))
            except OSError:
                continue
    return found


def _wait(condition, seconds=30.0):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        if condition():
            return True
        time.sleep(0.02)
    return False


def _start(args, env):
    return subprocess.Popen(
        [sys.executable, "-m", "keyline", *map(str, args)],
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        start_new_session=True,  # like a terminal: keyline leads its own process group
    )


def _send(proc, sig):
    if sig == signal.SIGINT:
        os.killpg(proc.pid, sig)  # a Ctrl-C reaches keyline's group, not the engine's
    else:
        os.kill(proc.pid, sig)  # `timeout` and `kill` signal the process


def _env(tmp_path, path):
    tmpdir = tmp_path / "tmpdir"
    tmpdir.mkdir()
    return dict(os.environ, PATH=str(path), TMPDIR=str(tmpdir)), tmpdir


def _no_keyline_dirs(tmpdir):
    return sorted(p.name for p in tmpdir.iterdir() if p.name.startswith("keyline-")) == []


# ---------------------------------------------------------------------------------------
# signals, with stand-in engines (runs anywhere, CI included)


@needs_sleep
@pytest.mark.parametrize("sig", [signal.SIGINT, signal.SIGTERM])
def test_a_signal_mid_validate_cleans_up(tmp_path, sig):
    pidfile = tmp_path / "child.pid"
    _stand_in(tmp_path / "bin", "officecli", pidfile, "1.0.152")
    env, tmpdir = _env(tmp_path, tmp_path / "bin")
    blocker = tmp_path / "no-render"
    blocker.touch()
    proc = _start(["check", CLEAN, "--mode", "read", "-o", blocker], env)
    assert _wait(pidfile.exists), proc.communicate(timeout=5)
    _send(proc, sig)
    _out, err = proc.communicate(timeout=30)
    assert proc.returncode == 128 + sig, err.decode()
    child = int(pidfile.read_text())
    assert _wait(lambda: not _alive(child), 5), "the officecli child survived"
    assert _no_keyline_dirs(tmpdir)
    assert b"internal error" not in err


@needs_sleep
@pytest.mark.parametrize("sig", [signal.SIGINT, signal.SIGTERM])
def test_a_signal_mid_render_cleans_up(tmp_path, sig):
    pytest.importorskip("pypdfium2")  # the rasterizer that makes auto pick LibreOffice
    pidfile = tmp_path / "child.pid"
    _stand_in(tmp_path / "bin", "soffice", pidfile, "LibreOffice 0.0 (stand-in)")
    env, tmpdir = _env(tmp_path, tmp_path / "bin")
    proc = _start(["render", KPI, "-o", tmp_path / "out"], env)
    assert _wait(pidfile.exists), proc.communicate(timeout=5)
    _send(proc, sig)
    _out, err = proc.communicate(timeout=30)
    assert proc.returncode == 128 + sig, err.decode()
    child = int(pidfile.read_text())
    assert _wait(lambda: not _alive(child), 5), "the soffice child survived"
    assert _no_keyline_dirs(tmpdir)


# ---------------------------------------------------------------------------------------
# signals, with the real engines


def _big_deck(path, slides=150):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Cm, Pt

    prs = pptx.Presentation()
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tf = s.shapes.add_textbox(Cm(2), Cm(2), Cm(20), Cm(10)).text_frame
        tf.text = f"Slide {i + 1}: " + "keepers wait a season " * 20
        tf.paragraphs[0].runs[0].font.size = Pt(18)
    prs.save(path)
    return path


@pytest.mark.skipif(render_mod.find_soffice() is None, reason="LibreOffice is not installed")
@pytest.mark.parametrize("sig", [signal.SIGINT, signal.SIGTERM])
def test_a_signal_mid_libreoffice_render_leaves_no_soffice(tmp_path, sig):
    deck = _big_deck(tmp_path / "big.pptx")
    env, tmpdir = _env(tmp_path, os.environ["PATH"])
    proc = _start(["render", deck, "-o", tmp_path / "out", "--engine", "libreoffice"], env)
    assert _wait(lambda: _mentioning(str(tmpdir)), 60), proc.communicate(timeout=5)
    time.sleep(0.5)  # let soffice start soffice.bin
    _send(proc, sig)
    _out, err = proc.communicate(timeout=60)
    assert proc.returncode == 128 + sig, err.decode()
    assert _wait(lambda: not _mentioning(str(tmpdir)), 10), _mentioning(str(tmpdir))
    assert _no_keyline_dirs(tmpdir)


@pytest.mark.officecli
@pytest.mark.parametrize("sig", [signal.SIGINT, signal.SIGTERM])
def test_a_signal_mid_officecli_validate_leaves_no_officecli(tmp_path, sig):
    env, tmpdir = _env(tmp_path, os.environ["PATH"])
    blocker = tmp_path / "no-render"
    blocker.touch()
    proc = _start(["check", CLEAN, "--mode", "read", "-o", blocker], env)
    if not _wait(lambda: _mentioning(str(tmpdir)), 60):
        proc.communicate(timeout=30)
        pytest.skip("officecli validate finished before a signal could reach it")
    _send(proc, sig)
    _out, err = proc.communicate(timeout=60)
    assert proc.returncode in (128 + sig, 0), err.decode()  # 0: it finished first
    assert _wait(lambda: not _mentioning(str(tmpdir)), 10), _mentioning(str(tmpdir))
    assert _no_keyline_dirs(tmpdir)


# ---------------------------------------------------------------------------------------
# odd output, a validate timeout, an empty -o


def test_engine_output_that_is_not_utf8(tmp_path):
    bin_dir = tmp_path / "bin"
    bin_dir.mkdir()
    exe = bin_dir / "officecli"
    exe.write_bytes(b"#!/bin/sh\nprintf 'Fehler: \\374berlauf\\n' >&2\nexit 1\n")
    exe.chmod(0o755)
    env = dict(os.environ, PATH=str(bin_dir))
    proc = keyline("check", CLEAN, "--mode", "read", "-o", tmp_path / "o", env=env)
    err = proc.stderr.decode()
    assert proc.returncode == 0 and "internal error" not in err
    assert "validate: skipped (officecli could not run: Fehler: �berlauf)" in err
    data = json.loads(keyline("doctor", "--json", env=env).stdout)
    (validator,) = [c for c in data["checks"] if c["check"] == "validator"]
    assert validator["token"] == "NO_VALIDATOR" and "�" in validator["detail"]


@needs_sleep
def test_a_validate_timeout_kills_the_group(tmp_path, monkeypatch):
    pidfile = tmp_path / "child.pid"
    _stand_in(tmp_path / "bin", "officecli", pidfile, "1.0.152")
    monkeypatch.setenv("PATH", str(tmp_path / "bin"))
    monkeypatch.setattr(validate_mod, "TIMEOUT_S", 1)
    (f,) = validate_mod.validate(CLEAN)
    assert f.message == "officecli validate timed out after 1 s"
    child = int(pidfile.read_text())
    assert _wait(lambda: not _alive(child), 5), "the native officecli was orphaned"


def test_an_empty_output_directory_is_unusable(tmp_path):
    planted = tmp_path / "slide-99.png"
    planted.write_bytes(b"x")
    proc = subprocess.run(
        [sys.executable, "-m", "keyline", "render", str(KPI), "-o", ""],
        cwd=tmp_path,
        capture_output=True,
    )
    assert proc.returncode == 1
    last = proc.stderr.decode().splitlines()[-1]
    assert last == "keyline: render failed: cannot use '' as the output directory: it is empty"
    assert planted.exists()
    check = keyline("check", KPI, "-o", "", "--no-validate")
    assert check.returncode == 2  # lint's exit code
    assert "render: skipped (cannot use '' as the output directory" in check.stderr.decode()
