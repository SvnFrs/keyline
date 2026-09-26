"""AC-5: python-pptx placeholder deck. Title and body geometry resolve from the layout,
with no adapter-unresolved for them. Expected values: python-pptx 1.0.2's default
template, re-derived in audit 01 and cross-checked against python-pptx below."""

import zipfile

import pytest

from keyline.lint import lint_path
from keyline.ooxml.adapter import load_deck
from tests.conftest import FOREIGN

DECK = FOREIGN / "pptx-default-placeholders.pptx"
EXPECTED = {
    "title": (457200, 274638, 8229600, 1143000),
    "obj": (457200, 1600200, 8229600, 4525963),
}


def _shapes():
    deck, _ = load_deck(DECK)
    return {s.ph_type: s for s in deck.slides[0].shapes}


def test_slide_has_no_xfrm():
    with zipfile.ZipFile(DECK) as z:
        assert b"a:xfrm" not in z.read("ppt/slides/slide1.xml")


def test_geometry_resolves_from_layout():
    shapes = _shapes()
    for ph_type, rect in EXPECTED.items():
        s = shapes[ph_type]
        assert (s.x, s.y, s.w, s.h) == rect, ph_type
        assert (s.box.x, s.box.y, s.box.w, s.box.h) == rect


def test_sizes_resolve_from_master_txstyles():
    shapes = _shapes()
    assert {r.size for r in shapes["title"].runs} == {4400}
    assert {r.size for r in shapes["obj"].runs} == {3200}


def test_background_resolves_through_bgref():
    deck, _ = load_deck(DECK)
    assert deck.slides[0].background == "solid:#FFFFFF"


def test_no_adapter_unresolved_for_the_placeholders():
    findings = lint_path(DECK).findings
    ids = {s.id for s in _shapes().values()}
    assert not [f for f in findings if f.rule == "adapter-unresolved" and f.shape_id in ids]


def test_python_pptx_oracle_agrees():
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation(DECK)
    got = {
        ph.placeholder_format.type: (ph.left, ph.top, ph.width, ph.height)
        for ph in prs.slides[0].placeholders
    }
    from pptx.enum.shapes import PP_PLACEHOLDER

    assert got[PP_PLACEHOLDER.TITLE] == EXPECTED["title"]
    assert got[PP_PLACEHOLDER.OBJECT] == EXPECTED["obj"]
