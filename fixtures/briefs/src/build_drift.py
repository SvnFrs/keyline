"""Build the AC-8 drift decks: python fixtures/briefs/src/build_drift.py [OUT_DIR]

base.pptx is what the pen will write for drift/base.brief.toml (spec 002 amendment B-3):
Swiss neutral layouts, text in the region placeholders, the keyline rule on evidence
slides, source lines, and the disclosure note on the last slide. Each drift is the base
with one pinned change (AC-8's table), so that exactly one finding moves. Output goes
through zipnorm, so rebuilding gives identical bytes.
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE.parent / "rules" / "src"))

import _deck as d  # noqa: E402

from keyline import zipnorm  # noqa: E402

OUT = HERE / "drift"
INK = "111111"
DISCLOSURE = "Note: BonsaiHub is a parody. Its trees and every number in this deck are fictional."


def base():
    prs = d.swiss_deck("neutral", "presented")
    d.swiss_slide(
        prs,
        "keyline:cover",
        notes=None,
        title="BonsaiHub",
        main="A slow dating app for trees and their keepers",
    )
    d.swiss_slide(
        prs,
        "keyline:statement",
        notes="Say it, then wait.",
        title="Every tree deserves a patient keeper",
    )
    s = d.swiss_slide(
        prs,
        "keyline:evidence",
        notes="Twelve thousand four hundred trees.",
        title="The waitlist grew faster than keepers could sign up",
        main="12,400 trees are waiting for a keeper",
        footer="Source: BonsaiHub waitlist, September 2026 (fictional)",
    )
    d.keyline_rule(s, INK)
    s = d.swiss_slide(
        prs,
        "keyline:evidence",
        notes="Seventy-one percent after ninety days.",
        title="Most keepers stay for the whole season",
        main="71% of keepers are still active after 90 days",
        footer="Source: BonsaiHub product analytics, September 2026 (fictional)",
    )
    d.keyline_rule(s, INK)
    d.swiss_slide(
        prs,
        "keyline:close",
        notes="End on the ask.",
        title="Put a keeper on every waiting tree",
        main="Join the waitlist as a keeper this season",
        footer=DISCLOSURE,
    )
    return prs


def _replace(slide, old, new):
    runs = [
        r
        for sh in slide.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        for r in p.runs
        if old in r.text
    ]
    assert len(runs) == 1, old
    runs[0].text = runs[0].text.replace(old, new)


def drift_headline(prs):  # edited, still within title_words_max
    _replace(prs.slides[2], "could sign up", "could join")


def drift_slide_count(prs):  # a copy of the close appended; the disclosure stays last
    close = prs.slides[4]
    copy_ = prs.slides.add_slide(close.slide_layout)
    for ph in list(copy_.placeholders):
        ph.element.getparent().remove(ph.element)
    for shape in close.shapes:
        copy_.shapes._spTree.append(copy.deepcopy(shape.element))
    copy_.notes_slide.notes_text_frame.text = close.notes_slide.notes_text_frame.text


def drift_unsourced(prs):  # in body text, not the headline
    _replace(prs.slides[2], "12,400", "12.4k")


def drift_source_missing(prs):  # the region stays occupied
    slide = prs.slides[3]
    _replace(
        slide,
        "Source: BonsaiHub product analytics, September 2026 (fictional)",
        "Note: figures are illustrative",
    )


def drift_undisclosed(prs):
    footer = [
        sh for sh in prs.slides[4].shapes if sh.has_text_frame and sh.text_frame.text == DISCLOSURE
    ]
    assert len(footer) == 1
    footer[0].element.getparent().remove(footer[0].element)


def drift_role(prs):  # a statement slide's layout swapped to keyline:quote
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    slide = prs.slides[1]
    quote = prs.slide_layouts.get_by_name("keyline:quote")
    (rel,) = [r for r in slide.part.rels.values() if r.reltype == RT.SLIDE_LAYOUT]
    rel._target = quote.part


DRIFTS = {
    "drift-headline": drift_headline,
    "drift-slide-count": drift_slide_count,
    "drift-unsourced": drift_unsourced,
    "drift-source-missing": drift_source_missing,
    "drift-undisclosed": drift_undisclosed,
    "drift-role": drift_role,
}


def _save(prs, out: Path, name: str) -> Path:
    buf = io.BytesIO()
    prs.save(buf)
    return zipnorm.write(out / f"{name}.pptx", zipnorm.read_entries(buf.getvalue()))


def build(out: Path = OUT) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    written = [_save(base(), out, "base")]
    for name, change in DRIFTS.items():
        prs = base()
        change(prs)
        written.append(_save(prs, out, name))
    return written


if __name__ == "__main__":
    for p in build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT):
        print(p)
