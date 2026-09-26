# Adapted from audit 02's stress-corpus/src/ppx_layouts.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D02: python-pptx default template, one slide per default layout (all 11).

Placeholders are filled with text and keep their inherited geometry (no xfrm on the
slide). Every slide has notes. Nothing is planted; whatever fires comes from the
default template itself.
"""
import sys

import _stress  # noqa: E402  (fixtures/foreign/stress/src)
from pptx import Presentation
from pptx.util import Pt

out = sys.argv[1]  # d02 is only an input to strict_convert.py; it is not committed
prs = Presentation()
L = {l.name: l for l in prs.slide_layouts}
PHOTO = _stress.photo()


def add(layout_name, texts, notes="Speaker notes here."):
    s = prs.slides.add_slide(L[layout_name])
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx in texts:
            val = texts[idx]
            if val == "PICTURE":
                ph.insert_picture(PHOTO)
            else:
                ph.text_frame.text = val
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


add("Title Slide", {0: "Annual planning", 1: "Priorities for the year ahead"})
add("Title and Content", {0: "Where we stand", 1: "Revenue grew in every region this year\nCosts stayed flat for the second year\nHiring is on plan"})
add("Section Header", {0: "Part two: the plan", 1: "What we will do next"})
add("Two Content", {0: "Before and after", 1: "Manual review of every order\nThree days to ship", 2: "Automated review of most orders\nOne day to ship"})
add("Comparison", {0: "Option A or option B", 1: "Option A", 2: "Cheaper to build and run\nSlower to ship", 3: "Option B", 4: "Faster to ship to customers\nCosts more to run"})
add("Title Only", {0: "A slide with only a title"})
add("Blank", {})
add("Content with Caption", {0: "Chart notes", 1: "The main content goes here in the big box", 2: "This caption explains the content to the left in a few more words"})
add("Picture with Caption", {0: "North warehouse", 1: "PICTURE", 2: "The site opened in August after a short delay"})
add("Title and Vertical Text", {0: "Vertical body", 1: "This body text runs vertically down the slide"})
add("Vertical Title and Text", {0: "Vertical title", 1: "This body text sits to the left of a vertical title"})
_stress.tyler(prs)
prs.save(out)
print("wrote", out)
