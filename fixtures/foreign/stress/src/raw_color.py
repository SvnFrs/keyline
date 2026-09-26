# Adapted from audit 02's stress-corpus/src/raw_color.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D16: color semantics that PowerPoint honors, written with python-pptx + lxml.

1 cover
2 p:clrMapOvr (tx1->lt1) on a slide with an explicit 111111 background; text has no
  color, so it resolves to tx1 = lt1 = white.                       render: white on near-black
3 p:clrMapOvr (bg1->dk1); background is bgRef bg1 (= dk1 = black); text is explicit
  222222.                                                          render: dark grey on black
4 layout-level dark panel: "Title Only" layout gets a navy rectangle over its left 45%
  and white title text in the layout lstStyle.                     render: white title on navy
5 alpha overlays: black fill at alpha 15% (renders ~D9D9D9) with black text (fine) and
  with white text (unreadable).
6 containment near-miss: white text box 0.1 in larger than the navy card under it on
  every side.                                                      render: white on navy
7 run alpha: white text at alpha 25% on navy.                     render: faint text
"""
import copy
import sys

import _stress  # noqa: E402  (fixtures/foreign/stress/src)
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

out = _stress.out_path("d16_raw_color.pptx")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
L = {l.name: l for l in prs.slide_layouts}
blank = L["Blank"]


def tb(s, x, y, w, h, text, size=24, color=None, bold=False):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.word_wrap = True
    r = t.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return t, r


def rect(s, x, y, w, h, fill_hex, alpha=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    sh.line.fill.background()
    # drop p:style so nothing but spPr decides the look
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    if alpha is not None:
        clr = sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha))
    return sh


def set_bg_xml(slide, xml):
    cSld = slide._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    cSld.insert(0, parse_xml(xml))


def clr_map_ovr(slide, **mapping):
    base = dict(bg1="lt1", tx1="dk1", bg2="lt2", tx2="dk2", accent1="accent1", accent2="accent2",
                accent3="accent3", accent4="accent4", accent5="accent5", accent6="accent6",
                hlink="hlink", folHlink="folHlink")
    base.update(mapping)
    sld = slide._element
    old = sld.find(qn("p:clrMapOvr"))
    if old is not None:
        sld.remove(old)
    ovr = etree.SubElement(sld, qn("p:clrMapOvr"))
    m = etree.SubElement(ovr, "{%s}overrideClrMapping" % A)
    for k, v in base.items():
        m.set(k, v)
    # p:clrMapOvr must come right after p:cSld
    sld.remove(ovr)
    sld.insert(list(sld).index(sld.find(qn("p:cSld"))) + 1, ovr)


def notes(s, t="n"):
    s.notes_slide.notes_text_frame.text = t


# 1 cover
s = prs.slides.add_slide(blank)
tb(s, 1, 3, 11.3, 1.2, "Color semantics", 44, "111111", True)
notes(s)

# 2 clrMapOvr with explicit dark background, text color left to tx1
s = prs.slides.add_slide(blank)
set_bg_xml(s, f'<p:bg xmlns:p="{P}" xmlns:a="{A}"><p:bgPr><a:solidFill><a:srgbClr val="111111"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>')
clr_map_ovr(s, bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2")
tb(s, 0.75, 0.6, 11.8, 1.0, "Dark slide via color map override", 40, None, True)
tb(s, 0.75, 2.0, 11.8, 2.0, "Body text with no explicit color uses tx1, which this slide maps to light", 24)
tb(s, 0.75, 4.5, 11.8, 2.0, "A second paragraph keeps the lower half of the slide from looking empty", 24)
notes(s)

# 3 clrMapOvr with bgRef bg1 (-> dk1 = black) and explicit dark grey text
s = prs.slides.add_slide(blank)
set_bg_xml(s, f'<p:bg xmlns:p="{P}" xmlns:a="{A}"><p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef></p:bg>')
clr_map_ovr(s, bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2")
tb(s, 0.75, 0.6, 11.8, 1.0, "Background maps bg1 to dark", 40, None, True)
tb(s, 0.75, 2.0, 11.8, 2.0, "This dark grey sentence sits on a black background and is hard to read", 24, "222222")
tb(s, 0.75, 4.5, 11.8, 2.0, "A second paragraph keeps the lower half of the slide from looking empty", 24)
notes(s)

# 4 layout-level dark panel: modify the "Title Only" layout
lay = L["Title Only"]
lsp = lay._element.find(qn("p:cSld")).find(qn("p:spTree"))
panel = parse_xml(
    f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="90" name="Navy panel"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{int(Inches(6.0))}" cy="{int(Inches(7.5))}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    f'<a:solidFill><a:srgbClr val="14213D"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
lsp.insert(2, panel)  # behind the placeholders
for sp in lsp.iter(qn("p:sp")):
    ph = sp.find(".//" + qn("p:ph"))
    if ph is not None and ph.get("type") == "title":
        spPr = sp.find(qn("p:spPr"))
        xfrm = spPr.find(qn("a:xfrm"))
        if xfrm is None:
            xfrm = parse_xml(f'<a:xfrm xmlns:a="{A}"><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>')
            spPr.insert(0, xfrm)
        xfrm.find(qn("a:off")).set("x", str(int(Inches(0.75))))
        xfrm.find(qn("a:off")).set("y", str(int(Inches(0.75))))
        xfrm.find(qn("a:ext")).set("cx", str(int(Inches(4.5))))
        xfrm.find(qn("a:ext")).set("cy", str(int(Inches(3.0))))
        body = sp.find(qn("p:txBody"))
        lst = body.find(qn("a:lstStyle"))
        lst.append(parse_xml(
            f'<a:lvl1pPr xmlns:a="{A}" algn="l"><a:defRPr sz="4000"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:defRPr></a:lvl1pPr>'))
s = prs.slides.add_slide(lay)
s.shapes.title.text = "Title on the layout panel"
tb(s, 6.75, 0.75, 5.8, 6.0, "The body text sits on the white right half of the slide, next to the navy panel", 24, "1A1A1A")
notes(s)

# 5 alpha overlays (black at 15% over white renders about D9D9D9)
s = prs.slides.add_slide(blank)
tb(s, 0.75, 0.6, 11.8, 1.0, "Translucent cards", 40, "111111", True)
rect(s, 0.75, 2.0, 5.6, 4.6, "000000", alpha=15000)
tb(s, 1.0, 2.3, 5.1, 4.0, "Black text on a pale translucent card reads fine", 24, "000000")
rect(s, 6.95, 2.0, 5.6, 4.6, "000000", alpha=15000)
tb(s, 7.2, 2.3, 5.1, 4.0, "White text on the same pale card cannot be read", 24, "FFFFFF")
notes(s)

# 6 containment near-miss
s = prs.slides.add_slide(blank)
tb(s, 0.75, 0.6, 11.8, 1.0, "Text box a bit larger than its card", 40, "111111", True)
rect(s, 1.0, 2.5, 11.3, 3.5, "1F3864")
tb(s, 0.9, 2.4, 11.5, 3.7, "White text centred on the navy card, in a box slightly larger than the card", 24, "FFFFFF")
notes(s)

# 7 run alpha: white at 25% on navy
s = prs.slides.add_slide(blank)
tb(s, 0.75, 0.6, 11.8, 1.0, "Faded text", 40, "111111", True)
rect(s, 0.75, 2.0, 11.8, 4.6, "1F3864")
t, r = tb(s, 1.0, 2.3, 11.3, 4.0, "This white sentence is drawn at twenty five percent opacity", 24, "FFFFFF")
clr = r._r.find(qn("a:rPr")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
etree.SubElement(clr, qn("a:alpha")).set("val", "25000")
notes(s)

_stress.tyler(prs)
prs.save(out)
print("wrote", out)
