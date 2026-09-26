# Adapted from audit 02's stress-corpus/src/ppx_localized_names.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D25: shape names as a localized PowerPoint writes them (Vietnamese, Japanese, German),
plus an emoji in a name. Each named shape carries a planted finding so its name is
printed. Used to test output encoding (cp1252 / ascii streams)."""
import _stress  # noqa: E402  (fixtures/foreign/stress/src)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Localized names"
s.shapes.title.name = "Tiêu đề 1"
s.notes_slide.notes_text_frame.text = "n"
for i, (name, y) in enumerate([("タイトル 2", 2.0), ("Textfeld 3 – Übersicht", 3.5), ("Box 🚀 4", 5.0)]):
    t = s.shapes.add_textbox(Inches(0.2), Inches(y), Inches(8), Inches(1))
    t.name = name
    r = t.text_frame.paragraphs[0].add_run(); r.text = "Grey text near the edge of the slide here"; r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
_stress.tyler(prs)
prs.save(_stress.out_path("d25_ppx_localized_names.pptx"))
print("ok")
