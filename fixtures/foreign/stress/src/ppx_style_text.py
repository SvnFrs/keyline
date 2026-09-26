# Adapted from audit 02's stress-corpus/src/ppx_style_text.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D15: python-pptx autoshapes whose text color comes only from p:style/a:fontRef.

python-pptx's add_shape() writes <p:style> with <a:fontRef idx="minor"><a:schemeClr
val="lt1"/></a:fontRef>, so text with no explicit color renders WHITE. The user then
sets a solid fill, which is the most common python-pptx pattern.

slide 2: dark navy fill, no text color set -> renders white on navy (fine).
slide 3: pale yellow fill, no text color set -> renders white on pale yellow (unreadable).
slide 4: explicit black text on the pale yellow fill (control: fine).
"""
import sys

import _stress  # noqa: E402  (fixtures/foreign/stress/src)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

out = _stress.out_path("d15_ppx_style_fontref.pptx")
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def title(s, text):
    t = s.shapes.add_textbox(Inches(0.75), Inches(0.6), Inches(11.8), Inches(1.0))
    r = t.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)


def card(s, fill_hex, text, color=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(2.0), Inches(11.8), Inches(4.6))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    sh.line.fill.background()
    sh.text_frame.text = text
    r = sh.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(24)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return sh


s = prs.slides.add_slide(blank)
title(s, "Style font colors")
s.notes_slide.notes_text_frame.text = "cover"

s = prs.slides.add_slide(blank)
title(s, "Navy card, default text color")
card(s, "1F3864", "This sentence has no color set so it takes the style font color")
s.notes_slide.notes_text_frame.text = "white on navy"

s = prs.slides.add_slide(blank)
title(s, "Pale card, default text color")
card(s, "FFF2CC", "This sentence has no color set so it takes the style font color")
s.notes_slide.notes_text_frame.text = "white on pale yellow"

s = prs.slides.add_slide(blank)
title(s, "Pale card, explicit black text")
card(s, "FFF2CC", "This sentence is explicitly black on the pale yellow card", color="000000")
s.notes_slide.notes_text_frame.text = "control"
_stress.tyler(prs)
prs.save(out)
print("wrote", out)
