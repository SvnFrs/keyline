# Adapted from audit 02's stress-corpus/src/lo_autofit.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D18: LibreOffice-written autofit. python-pptx writes a body placeholder whose text
overflows at the template size (32/28 pt) and marks it normAutofit ("shrink text on
overflow"); LibreOffice then computes the shrink and writes <a:normAutofit
fontScale=...> itself. Run lo_autofit.sh to produce the round-tripped deck.
"""
import sys

import _stress  # noqa: E402  (fixtures/foreign/stress/src)
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

prs = Presentation()
L = {l.name: l for l in prs.slide_layouts}
s = prs.slides.add_slide(L["Title Slide"])
s.shapes.title.text = "Autofit"
s.placeholders[1].text = "LibreOffice computes the shrink"
s.notes_slide.notes_text_frame.text = "cover"

s = prs.slides.add_slide(L["Title and Content"])
s.shapes.title.text = "Too much text"
body = s.placeholders[1]
lines = [
    "Revenue grew in every region this year, led by the north",
    "Costs stayed flat for the second year in a row",
    "Hiring is on plan for all teams except support",
    "Two new warehouses opened in the spring and summer",
    "Customer complaints fell for the third straight quarter",
    "Shipping times dropped by eleven percent overall",
    "The new carrier contract starts in the first quarter",
    "We will review the pricing model before the next budget",
    "Inventory turns improved in every warehouse we run",
    "Next year we plan to open one more site in the east",
]
body.text_frame.text = "\n".join(lines + lines)  # 20 lines
bodyPr = body.text_frame._txBody.find(qn("a:bodyPr"))
bodyPr.append(parse_xml('<a:normAutofit xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
s.notes_slide.notes_text_frame.text = "overflowing"
_stress.tyler(prs)
prs.save(sys.argv[1])  # input for LibreOffice; see lo_autofit.sh
print("ok")
