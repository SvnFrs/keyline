# Adapted from audit 02's stress-corpus/src/perf_shapes.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D22 perf decks: one content slide with N small filled text boxes on a grid (a data
map / pixel chart / dense diagram), plus a 60 x 150-shape deck. Built with python-pptx.

Grid cells are equal-sized, equally spaced, filled and each has text: the worst case for
equal-card-row and box-overlap (every row is a candidate).
"""
import math
import sys

import _stress  # noqa: E402  (fixtures/foreign/stress/src)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

# Usage: python perf_shapes.py OUT_DIR [all|perf_1000|perf_150_x60 ...]  (default: all)
import os
OUT = os.path.join(sys.argv[1] if len(sys.argv) > 1 else ".", "")
os.makedirs(OUT, exist_ok=True)
WANT = set(sys.argv[2:]) or {"all"}


def grid(slide, n, sw, sh, overlap=False):
    cols = math.ceil(math.sqrt(n * 16 / 9))
    rows = math.ceil(n / cols)
    x0, y0 = Inches(0.6), Inches(1.8)
    cw = (sw - 2 * x0) // cols
    ch = (sh - y0 - Inches(0.6)) // rows
    k = 0
    for r in range(rows):
        for c in range(cols):
            if k >= n:
                return
            w = cw * (1.6 if overlap else 0.8)
            h = ch * (1.6 if overlap else 0.8)
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 + c * cw, y0 + r * ch, int(w), int(h))
            s.fill.solid(); s.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            s.text_frame.text = str(k)
            s.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
            k += 1


def deck(n, overlap=False, slides=1):
    stem = f"perf_{n}{'_overlap' if overlap else ''}{'_x' + str(slides) if slides > 1 else ''}"
    if "all" not in WANT and stem not in WANT:
        return
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"{n} shapes"
        grid(s, n, prs.slide_width, prs.slide_height, overlap)
        s.notes_slide.notes_text_frame.text = "n"
    name = f"perf_{n}{'_overlap' if overlap else ''}{'_x' + str(slides) if slides > 1 else ''}.pptx"
    _stress.tyler(prs)
    prs.save(OUT + name)
    print("wrote", name)


for n in (100, 250, 500, 1000, 2000):
    deck(n)
deck(500, overlap=True)
deck(1000, overlap=True)
deck(150, slides=60)
