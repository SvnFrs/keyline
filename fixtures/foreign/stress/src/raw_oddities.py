# Adapted from audit 02's stress-corpus/src/raw_oddities.py: output paths, test photo and the Tyler author only (see _stress.py).
"""D24 crash hunt: one file per schema oddity, derived from a small python-pptx base.

Each case edits one part with a string function. The goal is only: no traceback, and
an exit code in {0, 1, 2}. Cases marked (valid) are schema-valid or tolerated by
PowerPoint/LibreOffice; for those, exit 1 would be a refusal of a legitimate deck.
"""
import os
import re
import sys
import tempfile
import zipfile

import _stress  # noqa: E402  (fixtures/foreign/stress/src)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

# Usage: python raw_oddities.py [OUT_DIR] [case ...]  (default: fixtures/foreign/stress)
ARGS = sys.argv[1:]
OUT = os.path.join(ARGS.pop(0) if ARGS and os.path.isdir(ARGS[0]) else str(_stress.OUT_DIR), "")
BASE = os.path.join(tempfile.mkdtemp(), "odd_base.pptx")
# Only the cases audit 02's acceptance criteria use are built by default.
KEEP = ARGS or ["o51_lummod_percent_string", "o52_alpha_percent_string",
    "o53_tint_percent_in_master_bg", "o54_lumoff_float", "o55_cxn_id_word", "o56_ext_40_digits"]

prs = Presentation()
L = {l.name: l for l in prs.slide_layouts}
s = prs.slides.add_slide(L["Title Slide"]); s.shapes.title.text = "Base"; s.placeholders[1].text = "cover"
s.notes_slide.notes_text_frame.text = "n"
s = prs.slides.add_slide(L["Title Only"]); s.shapes.title.text = "Content"
tb = s.shapes.add_textbox(Inches(0.75), Inches(2), Inches(8.5), Inches(1.5))
r = tb.text_frame.paragraphs[0].add_run(); r.text = "A paragraph of body text with more than five words in it"; r.font.size = Pt(24)
r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
g = s.shapes.add_group_shape()
a = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(2), Inches(1)); a.text_frame.text = "In group"
b = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(4), Inches(2), Inches(1)); b.text_frame.text = "Also in group"
c = s.shapes.add_connector(1, Inches(3), Inches(4.5), Inches(4), Inches(4.5))
c.begin_connect(a, 3); c.end_connect(b, 1)
s.notes_slide.notes_text_frame.text = "n"
_stress.tyler(prs)
prs.save(BASE)

SLIDE = "ppt/slides/slide2.xml"
PRES = "ppt/presentation.xml"
MASTER = "ppt/slideMasters/slideMaster1.xml"
THEME = "ppt/theme/theme1.xml"


def first_run_color(xml, new):
    return xml.replace('<a:srgbClr val="202020"/>', new, 1)


def sub1(pat, rep):
    return lambda x: re.sub(pat, rep, x, count=1, flags=re.S)


CASES = {
    # geometry
    "o01_xfrm_no_off": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:xfrm>)<a:off [^>]*/>', r"\1")),
    "o02_xfrm_no_ext": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:xfrm>)(<a:off [^>]*/>)<a:ext [^>]*/>', r"\1\2")),
    "o03_sp_no_xfrm_valid": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?)<a:xfrm>.*?</a:xfrm>', r"\1")),
    "o04_grp_no_choff_chext_valid": (SLIDE, sub1(r"<a:chOff [^>]*/><a:chExt [^>]*/>", "")),
    "o05_grp_chext_zero": (SLIDE, sub1(r'<a:chExt cx="\d+" cy="\d+"/>', '<a:chExt cx="0" cy="0"/>')),
    "o06_grp_no_xfrm_valid": (SLIDE, sub1(r"<p:grpSpPr><a:xfrm>.*?</a:xfrm></p:grpSpPr>", "<p:grpSpPr/>")),
    "o07_rot_negative": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?)<a:xfrm>', r'\1<a:xfrm rot="-5400000">')),
    "o08_rot_over_360": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?)<a:xfrm>', r'\1<a:xfrm rot="43200000">')),
    "o09_coord_float_text": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:off x=")\d+', r"\g<1>1.5e6")),
    "o10_coord_huge": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:ext cx=")\d+', r"\g<1>9223372036854775807")),
    "o11_ext_negative": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:ext cx=")\d+', r"\g<1>-914400")),
    "o12_sldsz_missing": (PRES, sub1(r"<p:sldSz [^>]*/>", "")),
    "o13_sldsz_zero": (PRES, sub1(r"<p:sldSz [^>]*/>", '<p:sldSz cx="0" cy="0"/>')),
    "o14_graphicframe_no_xfrm": (SLIDE, lambda x: x.replace("</p:spTree>",
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="90" name="GF"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table"><a:tbl><a:tblGrid/></a:tbl></a:graphicData></a:graphic></p:graphicFrame></p:spTree>')),
    "o15_pic_no_blip": (SLIDE, lambda x: x.replace("</p:spTree>",
        '<p:pic><p:nvPicPr><p:cNvPr id="91" name="Pic"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill/>'
        '<p:spPr><a:xfrm><a:off x="100" y="100"/><a:ext cx="100" cy="100"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic></p:spTree>')),
    "o16_cxn_dangling_ids": (SLIDE, lambda x: re.sub(r'<a:stCxn id="\d+"', '<a:stCxn id="9999"', x)),
    "o17_contentpart_ink": (SLIDE, lambda x: x.replace("</p:spTree>",
        '<p:contentPart xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" r:id="rId99"/></p:spTree>')),
    "o18_orphan_placeholder": (SLIDE, lambda x: x.replace("</p:spTree>",
        '<p:sp><p:nvSpPr><p:cNvPr id="92" name="Subtitle 9"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="subTitle" idx="7"/></p:nvPr></p:nvSpPr>'
        '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>Orphan placeholder with some words in it</a:t></a:r></a:p></p:txBody></p:sp></p:spTree>')),
    "o19_nested_groups_40_valid": (SLIDE, lambda x: x.replace("</p:spTree>",
        "".join(f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{200+i}" name="G{i}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="914400" y="914400"/><a:ext cx="914400" cy="914400"/><a:chOff x="0" y="0"/><a:chExt cx="914400" cy="914400"/></a:xfrm></p:grpSpPr>' for i in range(40))
        + '<p:sp><p:nvSpPr><p:cNvPr id="300" name="Deep"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>deep</a:t></a:r></a:p></p:txBody></p:sp>'
        + "</p:grpSp>" * 40 + "</p:spTree>")),
    "o20_nested_groups_300_valid": (SLIDE, lambda x: x.replace("</p:spTree>",
        "".join(f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{400+i}" name="G{i}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>' for i in range(300))
        + '<p:sp><p:nvSpPr><p:cNvPr id="999" name="Deep"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>deep</a:t></a:r></a:p></p:txBody></p:sp>'
        + "</p:grpSp>" * 300 + "</p:spTree>")),
    # text
    "o21_sz_zero": (SLIDE, lambda x: x.replace('sz="2400"', 'sz="0"', 1)),
    "o22_sz_huge": (SLIDE, lambda x: x.replace('sz="2400"', 'sz="400000"', 1)),
    "o23_sz_not_int": (SLIDE, lambda x: x.replace('sz="2400"', 'sz="24pt"', 1)),
    "o24_txbody_no_p": (SLIDE, sub1(r"(<p:cNvPr id=\"3\" .*?<a:lstStyle/>)<a:p>.*?</a:p>", r"\1")),
    "o25_field_and_br_only": (SLIDE, sub1(r"(<p:cNvPr id=\"3\" .*?<a:lstStyle/>)<a:p>.*?</a:p>",
        r'\1<a:p><a:fld id="{1}" type="slidenum"><a:rPr lang="en-US" sz="1200"/><a:t>2</a:t></a:fld><a:br><a:rPr/></a:br></a:p>')),
    "o26_latin_theme_ea_ref": (SLIDE, lambda x: x.replace('<a:rPr lang="en-US" sz="2400"', '<a:rPr lang="en-US" sz="2400"', 1).replace(
        '<a:srgbClr val="202020"/></a:solidFill>', '<a:srgbClr val="202020"/></a:solidFill><a:latin typeface="+mj-ea"/>', 1)),
    "o27_empty_typeface": (SLIDE, lambda x: x.replace('<a:srgbClr val="202020"/></a:solidFill>', '<a:srgbClr val="202020"/></a:solidFill><a:latin typeface=""/>', 1)),
    "o28_lvl_out_of_range": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:p>)', r'\1<a:pPr lvl="12"/>')),
    # color
    "o29_hslclr": (SLIDE, lambda x: first_run_color(x, '<a:hslClr hue="14400000" sat="100000" lum="50000"/>')),
    "o30_prstclr": (SLIDE, lambda x: first_run_color(x, '<a:prstClr val="lightGray"/>')),
    "o31_sysclr_no_lastclr": (SLIDE, lambda x: first_run_color(x, '<a:sysClr val="windowText"/>')),
    "o32_scrgbclr": (SLIDE, lambda x: first_run_color(x, '<a:scrgbClr r="50000" g="50000" b="50000"/>')),
    "o33_transforms_exotic": (SLIDE, lambda x: first_run_color(x, '<a:srgbClr val="202020"><a:hueOff val="100"/><a:satMod val="200000"/><a:inv/><a:comp/><a:gray/><a:gamma/><a:invGamma/></a:srgbClr>')),
    "o34_lummod_over": (SLIDE, lambda x: first_run_color(x, '<a:schemeClr val="tx1"><a:lumMod val="500000"/><a:lumOff val="-300000"/></a:schemeClr>')),
    "o35_phclr_bare": (SLIDE, lambda x: first_run_color(x, '<a:schemeClr val="phClr"/>')),
    "o36_srgb_bad_hex": (SLIDE, lambda x: first_run_color(x, '<a:srgbClr val="GGHHII"/>')),
    "o37_srgb_short_hex": (SLIDE, lambda x: first_run_color(x, '<a:srgbClr val="FFF"/>')),
    # master / theme
    "o38_master_no_clrmap": (MASTER, sub1(r"<p:clrMap [^>]*/>", "")),
    "o39_master_no_txstyles_valid": (MASTER, sub1(r"<p:txStyles>.*</p:txStyles>", "")),
    "o40_theme_no_fontscheme": (THEME, sub1(r"<a:fontScheme .*?</a:fontScheme>", "")),
    "o41_theme_no_clrscheme": (THEME, sub1(r"<a:clrScheme .*?</a:clrScheme>", "")),
    "o42_bgref_idx_huge": (MASTER, lambda x: x.replace('<p:bgRef idx="1001">', '<p:bgRef idx="1999">')),
    "o43_extlst_everywhere_valid": (SLIDE, lambda x: x.replace("</p:cNvPr>", "</p:cNvPr>").replace(
        '<p:cNvPr id="3" name="TextBox 2"/>',
        '<p:cNvPr id="3" name="TextBox 2"><a:extLst><a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}"><a16:creationId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" id="{00000000-0000-0000-0000-000000000001}"/></a:ext></a:extLst></p:cNvPr>')),
    "o44_ac_in_group_valid": (SLIDE, lambda x: x.replace("</p:grpSp>",
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">'
        '<p:sp><p:nvSpPr><p:cNvPr id="93" name="AC child"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="914400" y="4000000"/><a:ext cx="914400" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>ac</a:t></a:r></a:p></p:txBody></p:sp>'
        '</mc:Choice></mc:AlternateContent></p:grpSp>', 1)),
    "o45_unknown_element_in_sptree": (SLIDE, lambda x: x.replace("</p:spTree>", '<foo:bar xmlns:foo="urn:x"><foo:baz/></foo:bar></p:spTree>')),
    "o46_shape_id_duplicate_valid": (SLIDE, lambda x: x.replace('<p:cNvPr id="3"', '<p:cNvPr id="2"', 1)),
    "o47_shape_id_missing": (SLIDE, lambda x: x.replace('<p:cNvPr id="3" ', '<p:cNvPr ', 1)),
    "o48_shape_id_text": (SLIDE, lambda x: x.replace('<p:cNvPr id="3"', '<p:cNvPr id="three"', 1)),
    "o49_notes_rel_to_slide": ("ppt/slides/_rels/slide2.xml.rels", lambda x: x.replace("../notesSlides/notesSlide2.xml", "../slides/slide1.xml")),
    "o50_layout_rel_to_master": ("ppt/slides/_rels/slide2.xml.rels", lambda x: x.replace("../slideLayouts/slideLayout6.xml", "../slideMasters/slideMaster1.xml")),
    # values written the way ISO/IEC 29500 Strict writes percentages ("75%"), and other
    # non-integer numbers in integer attributes
    "o51_lummod_percent_string": (SLIDE, lambda x: first_run_color(x, '<a:schemeClr val="tx1"><a:lumMod val="75%"/><a:lumOff val="25%"/></a:schemeClr>')),
    "o52_alpha_percent_string": (SLIDE, lambda x: first_run_color(x, '<a:srgbClr val="202020"><a:alpha val="50%"/></a:srgbClr>')),
    "o53_tint_percent_in_master_bg": (MASTER, lambda x: x.replace('<p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef>',
        '<p:bgPr><a:solidFill><a:schemeClr val="bg1"><a:tint val="95%"/></a:schemeClr></a:solidFill><a:effectLst/></p:bgPr>')),
    "o54_lumoff_float": (SLIDE, lambda x: first_run_color(x, '<a:schemeClr val="tx1"><a:lumMod val="75000"/><a:lumOff val="25000.0"/></a:schemeClr>')),
    "o55_cxn_id_word": (SLIDE, lambda x: re.sub(r'<a:stCxn id="\d+"', '<a:stCxn id="first"', x)),
    "o56_ext_40_digits": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:ext cx=")\d+', r"\g<1>" + "9" * 40)),
    "o57_rotated_ext_25_digits": (SLIDE, lambda x: sub1(r'(<p:cNvPr id="3" .*?<a:ext cx=")\d+', r"\g<1>" + "9" * 25)(x).replace(
        '<p:cNvPr id="3" name="TextBox 2"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>',
        '<p:cNvPr id="3" name="TextBox 2"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm rot="2700000">')),
    "o58_off_universal_measure": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?<a:off x=")\d+', r"\g<1>2cm")),
    "o59_sz_float": (SLIDE, lambda x: x.replace('sz="2400"', 'sz="2400.5"', 1)),
    "o60_rot_float": (SLIDE, sub1(r'(<p:cNvPr id="3" .*?)<a:xfrm>', r'\1<a:xfrm rot="5400000.5">')),
}


def build(name, part, fn):
    zin = zipfile.ZipFile(BASE)
    with zipfile.ZipFile(OUT + name + ".pptx", "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zin.infolist():
            d = zin.read(info.filename)
            if info.filename == part:
                before = d.decode("utf-8")
                after = fn(before)
                if after == before:
                    print("WARNING: mutation had no effect:", name)
                d = after.encode("utf-8")
            zo.writestr(info, d)


for name in KEEP:
    part, fn = CASES[name]
    build(name, part, fn)
print(len(KEEP), "cases")
