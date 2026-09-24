"""Build the foreign fixtures: python fixtures/foreign/src/build_foreign.py [OUT_DIR]

pptx-default-placeholders.pptx (AC-5)
    python-pptx 1.0.2's default template (4:3), layout 1 "Title and Content". The title
    and body placeholders carry text and no a:xfrm, so their geometry must come from the
    layout.

nested-groups.pptx (AC-6)
    Raw p:grpSp XML with known absolute geometry, worked out by hand (EMU, 1 cm = 360000):

    G1 (top level): off (360000, 360000), ext (3600000, 1800000),
                    chOff (0, 0), chExt (7200000, 3600000)            -> scale 0.5
      G2 (in G1):   off (720000, 720000), ext (2880000, 1440000),
                    chOff (360000, 360000), chExt (1440000, 720000)   -> scale 2
        A (in G2):  off (360000, 360000) ext (720000, 360000)
                    G1 space: x = 720000 + (360000 - 360000)*2 = 720000, y = 720000,
                              w = 1440000, h = 720000
                    slide:    x = 360000 + 720000*0.5 = 720000, y = 720000,
                              w = 720000, h = 360000
        B (in G2):  off (1080000, 720000) ext (360000, 360000)
                    G1 space: x = 720000 + 720000*2 = 2160000, y = 720000 + 360000*2 = 1440000,
                              w = h = 720000
                    slide:    x = 360000 + 1080000 = 1440000, y = 360000 + 720000 = 1080000,
                              w = h = 360000
      C (in G1):    off (4320000, 2160000) ext (1440000, 720000)
                    slide:    x = 360000 + 2160000 = 2520000, y = 360000 + 1080000 = 1440000,
                              w = 720000, h = 360000
      G3 (in G1):   off (5040000, 360000), ext = chExt = (1440000, 720000), chOff (0, 0),
                    rot 90°                                            -> scale 1, turned
        D (in G3):  off (0, 0) ext (720000, 360000); center (360000, 180000)
                    G1 space before turning: center (5400000, 540000);
                    G3 center (5760000, 720000); offset (-360000, -180000) turned 90°
                    clockwise (y down) -> (180000, -360000); center (5940000, 360000),
                    rot 90°, w 720000, h 360000
                    slide: center (360000 + 2970000, 360000 + 180000) = (3330000, 540000),
                           w 360000, h 180000, rot 90°
                           unrotated rect x = 3150000, y = 450000
                           bounding box x = 3240000, y = 360000, w = 180000, h = 360000
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "rules" / "src"))

import _deck as d

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def default_placeholders():
    from pptx import Presentation

    prs = Presentation()  # the default template, 4:3, left as it is
    cp = prs.core_properties
    cp.author = cp.last_modified_by = "Tyler"
    cp.revision = 1
    cp.created = cp.modified = d.FIXED_TIME
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Placeholder title from the layout"
    s.placeholders[1].text = "Body text that inherits its box and size from the layout"
    return prs


def _grp(sid, name, off, ext, ch_off, ch_ext, rot=0, children=""):
    r = f' rot="{rot}"' if rot else ""
    return (
        f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvGrpSpPr/><p:nvPr/>'
        f'</p:nvGrpSpPr><p:grpSpPr><a:xfrm{r}><a:off x="{off[0]}" y="{off[1]}"/>'
        f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/><a:chOff x="{ch_off[0]}" y="{ch_off[1]}"/>'
        f'<a:chExt cx="{ch_ext[0]}" cy="{ch_ext[1]}"/></a:xfrm></p:grpSpPr>{children}</p:grpSp>'
    )


def _leaf(sid, name, off, ext, color="1E2761"):
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" '
        f'cy="{ext[1]}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></p:spPr></p:sp>'
    )


def nested_groups():
    prs = d.new_deck()
    s = d.slide(prs, bg="FFFFFF", notes="Groups with known absolute geometry.")
    g2 = _grp(
        20,
        "G2",
        (720000, 720000),
        (2880000, 1440000),
        (360000, 360000),
        (1440000, 720000),
        children=_leaf(21, "A", (360000, 360000), (720000, 360000))
        + _leaf(22, "B", (1080000, 720000), (360000, 360000)),
    )
    g3 = _grp(
        30,
        "G3",
        (5040000, 360000),
        (1440000, 720000),
        (0, 0),
        (1440000, 720000),
        rot=5400000,
        children=_leaf(31, "D", (0, 0), (720000, 360000)),
    )
    g1 = _grp(
        10,
        "G1",
        (360000, 360000),
        (3600000, 1800000),
        (0, 0),
        (7200000, 3600000),
        children=g2 + _leaf(11, "C", (4320000, 2160000), (1440000, 720000)) + g3,
    )
    wrapped = etree.fromstring(f'<w xmlns:a="{A}" xmlns:p="{P}">{g1}</w>')
    s.shapes._spTree.append(wrapped[0])
    return prs


DECKS = {
    "pptx-default-placeholders": default_placeholders,
    "nested-groups": nested_groups,
}


def build(out_dir: Path, names: list[str] | None = None) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for name in names or sorted(DECKS):
        path = out_dir / f"{name}.pptx"
        DECKS[name]().save(path)
        written.append(path)
    return written


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1]
    for p in build(out):
        print(p)
