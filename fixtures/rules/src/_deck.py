"""python-pptx helpers for building rule fixtures. Dev-only: never imported by keyline.

Positions and sizes are in cm; colors are "RRGGBB". Every deck is 16:9 (33.867 × 19.05 cm)
and carries the project's identity in its core properties.
"""

from __future__ import annotations

import datetime as dt

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

BLANK_LAYOUT = 6
FIXED_TIME = dt.datetime(2026, 9, 24, 0, 0, 0)
ALIGN = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    cp = prs.core_properties
    cp.author = cp.last_modified_by = "Tyler"
    cp.title = ""
    cp.revision = 1
    cp.created = cp.modified = FIXED_TIME
    return prs


def slide(prs, bg: str | None = None, notes: str | None = None):
    s = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    if bg is not None:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    if notes is not None:
        s.notes_slide.notes_text_frame.text = notes
    return s


def _style_runs(tf, size, bold, color, font, align):
    for p in tf.paragraphs:
        if align:
            p.alignment = ALIGN[align]
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = RGBColor.from_string(color)


def text(
    s,
    x,
    y,
    w,
    h,
    body,
    size=18,
    *,
    name=None,
    bold=False,
    color="111111",
    font="Arial",
    align=None,
    rot=0,
):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = body.split("\n")
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    _style_runs(tf, size, bold, color, font, align)
    if name:
        tb.name = name
    if rot:
        tb.rotation = rot
    return tb


def rect(
    s,
    x,
    y,
    w,
    h,
    fill="1E2761",
    *,
    name=None,
    body=None,
    size=18,
    color="FFFFFF",
    font="Arial",
    bold=False,
    shape=MSO_SHAPE.RECTANGLE,
):
    sh = s.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background()
    if body is not None:
        sh.text_frame.text = body
        _style_runs(sh.text_frame, size, bold, color, font, "ctr")
    if name:
        sh.name = name
    return sh


def connector(s, a, b):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(a, 3)
    c.end_connect(b, 1)
    return c


def picture_placeholder(s, x, y, w, h, name="picture"):
    """A tiny PNG picture (1×1 px, generated in memory)."""
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (1, 1), (200, 200, 200)).save(buf, format="PNG")
    buf.seek(0)
    pic = s.shapes.add_picture(buf, Cm(x), Cm(y), Cm(w), Cm(h))
    pic.name = name
    return pic
