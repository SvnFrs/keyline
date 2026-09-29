"""Measure LibreOffice's line pitch and wrap margin (spec 002 amendment B-7, task T-20).

    python tools/measure_lo.py [OUT_FILE]

Dev only (python-pptx, fontTools, pypdfium2). For each portable family (thresholds.toml
`portable_fonts`) the probes name the family itself, not its metric twin, so fontconfig
substitutes exactly as it does for a pen deck (audit 03 on Q-44). The script first prints
`fc-match` for every family and stops if one does not resolve to its twin.

The probe decks are converted with keyline's own render path (`render.convert_to_pdf`,
the same soffice command), and glyph origins are read back from the PDF:
- line pitch: four lines in one paragraph (line breaks) and four one-line paragraphs, at
  13, 24, 48 and 60 pt, regular and bold, line spacing 100 %, no paragraph spacing;
- wrap margin: a bold 14 pt label in a box 1.000×, 1.002×, 1.005× and 1.010× its advance
  sum (fontTools, no kerning), zero insets, wrap on, no autofit; the line count is read
  back.

§6.4's LibreOffice 24.2 values are the reference: a pitch of 1.20 × size (Liberation Sans),
and a bold 14 pt label wraps at 1.000× and 1.002× but not at 1.005× its advance sum.
Every number printed is tagged with the `soffice --version` that produced it.
"""

from __future__ import annotations

import ctypes
import datetime as dt
import io
import itertools
import platform
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from keyline import __version__  # noqa: E402
from keyline.config import load as load_config  # noqa: E402
from keyline.render import convert_to_pdf, find_soffice, libreoffice_version  # noqa: E402

EXPECTED_VERSION = "26.8.0.3"
SIZES = (13, 24, 48, 60)
WEIGHTS = ("regular", "bold")
FACTORS = ("1.000", "1.002", "1.005", "1.010")
SCAN = tuple(f"{0.980 + i * 0.0005:.4f}" for i in range(61))  # 0.9800 … 1.0100
LABELS = ("ACTIVE KEEPERS BY MONTH", "median wait for a match", "Oldest tree on the waitlist")
LINE = "Hxgh Ápq"
PITCH_EM = 1.2
PITCH_TOLERANCE = 0.005  # LibreOffice positions in 1/100 mm; 0.005 em is 0.07 pt at 13 pt
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
EMU_PER_PT = 12700


def fc_match(pattern: str, fmt: str) -> str:
    out = subprocess.run(["fc-match", "-f", fmt, pattern], capture_output=True, text=True)
    return out.stdout.split(",")[0].strip()


def check_fonts(families) -> list[str]:
    lines, bad = [], []
    for family, twin in families:
        got = fc_match(family, "%{family}")
        ok = got.casefold() == twin.casefold()
        lines.append(f"  {family:<16} -> {got:<18} twin {twin:<18} {'ok' if ok else 'MISMATCH'}")
        if not ok:
            bad.append(family)
    if bad:
        print("\n".join(lines))
        sys.exit(f"stop: {', '.join(bad)} do not resolve to their metric twins (audit 03, Q-44)")
    return lines


def advance_sum(family: str, text: str, size: float) -> float:
    from fontTools.ttLib import TTFont

    path = fc_match(f"{family}:bold", "%{file}")
    font = TTFont(path)
    cmap, hmtx = font.getBestCmap(), font["hmtx"]
    upm = font["head"].unitsPerEm
    total = sum(hmtx[cmap[ord(ch)]][0] for ch in text)
    return total * size / upm


def _box(slide, x_pt, y_pt, w_pt, h_pt):
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu

    box = slide.shapes.add_textbox(
        Emu(round(x_pt * EMU_PER_PT)),
        Emu(round(y_pt * EMU_PER_PT)),
        Emu(round(w_pt * EMU_PER_PT)),
        Emu(round(h_pt * EMU_PER_PT)),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _style(paragraph, family, size, bold):
    from pptx.util import Pt

    paragraph.line_spacing = 1.0
    paragraph.space_before = paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = family
        run.font.size = Pt(size)
        run.font.bold = bold


def build_scan(families):
    """The wrap threshold per label: one slide per factor in SCAN."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
    probes = []
    for family, _twin in families:
        for label in LABELS:
            width = advance_sum(family, label, 14)
            for factor in SCAN:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                tf = _box(slide, 72, 36.0, width * float(factor), 200)
                tf.paragraphs[0].text = label
                _style(tf.paragraphs[0], family, 14, True)
                probes.append((family, label, factor))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), probes


def build_probes(families):
    """(deck bytes, probe list); probe i is on slide i + 1."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
    probes = []
    top = 36.0  # pt
    for family, twin in families:
        for weight in WEIGHTS:
            for size in SIZES:
                for kind in ("breaks", "paragraphs"):
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    tf = _box(slide, 72, top, 800, 480)
                    if kind == "breaks":
                        tf.paragraphs[0].text = "\v".join([LINE] * 4)  # \v is <a:br/>
                    else:
                        tf.text = "\n".join([LINE] * 4)
                    for p in tf.paragraphs:
                        _style(p, family, size, weight == "bold")
                    probes.append(
                        {
                            "what": "pitch",
                            "family": family,
                            "twin": twin,
                            "weight": weight,
                            "size": size,
                            "kind": kind,
                            "top": top,
                        }
                    )
        for label in LABELS:
            width = advance_sum(family, label, 14)
            for factor in FACTORS:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                box_w = width * float(factor)
                tf = _box(slide, 72, top, box_w, 200)
                tf.paragraphs[0].text = label
                _style(tf.paragraphs[0], family, 14, True)
                probes.append(
                    {
                        "what": "wrap",
                        "family": family,
                        "twin": twin,
                        "label": label,
                        "advance": width,
                        "factor": factor,
                        "box": box_w,
                    }
                )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), probes


def read_pages(pdf: Path):
    """Per page: (baselines top to bottom, font names, page height), from glyph origins."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    doc = pdfium.PdfDocument(str(pdf))
    pages = []
    try:
        for page in doc:
            height = page.get_height()
            tp = page.get_textpage()
            ys, fonts = set(), set()
            x, y = ctypes.c_double(), ctypes.c_double()
            buf = ctypes.create_string_buffer(256)
            flags = ctypes.c_int()
            for i in range(tp.count_chars()):
                if not tp.get_text_range(i, 1).strip():
                    continue
                raw.FPDFText_GetCharOrigin(tp.raw, i, ctypes.byref(x), ctypes.byref(y))
                ys.add(round(y.value, 3))
                n = raw.FPDFText_GetFontInfo(tp.raw, i, buf, 256, ctypes.byref(flags))
                fonts.add(buf.raw[: max(n - 1, 0)].decode("utf-8", "replace"))
            pages.append((sorted(ys, reverse=True), sorted(fonts), height))
            tp.close()
    finally:
        doc.close()
    return pages


def main(argv: list[str]) -> int:
    families = load_config().portable_fonts
    soffice = find_soffice()
    if soffice is None:
        sys.exit("stop: soffice is not installed")
    version = libreoffice_version(soffice)
    out = [
        "# LibreOffice measurements for spec 002 amendment B-7 (task T-20)",
        f"version: {version}",
        f"date: {dt.date.today().isoformat()}  keyline {__version__}  {platform.platform()}",
    ]
    if EXPECTED_VERSION not in version:
        out.append(f"note: expected LibreOffice {EXPECTED_VERSION}; this run used {version}")
    out += ["", "fc-match (probes name the portable family; fontconfig picks the twin):"]
    out += check_fonts(families)

    deck_bytes, probes = build_probes(families)
    with tempfile.TemporaryDirectory(prefix="keyline-measure-") as tmp:
        deck = Path(tmp) / "probes.pptx"
        deck.write_bytes(deck_bytes)
        work = Path(tmp) / "work"
        work.mkdir()
        pdf = convert_to_pdf(deck, soffice, work)
        pages = read_pages(pdf)
        scan_bytes, scan_probes = build_scan(families)
        scan_deck = Path(tmp) / "scan.pptx"
        scan_deck.write_bytes(scan_bytes)
        work2 = Path(tmp) / "work2"
        work2.mkdir()
        scan_pages = read_pages(convert_to_pdf(scan_deck, soffice, work2))
    if len(pages) != len(probes):
        sys.exit(f"stop: {len(pages)} PDF pages for {len(probes)} probes")

    out += ["", f"line pitch (em = pitch / size), {version}:"]
    out.append(
        f"  {'family':<16} {'weight':<8} {'size':>4}  {'breaks':>7} {'paras':>7}"
        f"  {'1st baseline':>12}  embedded font"
    )
    pitch = {}
    rows = {}
    for probe, (ys, fonts, height) in zip(probes, pages, strict=True):
        if probe["what"] != "pitch":
            continue
        steps = [a - b for a, b in itertools.pairwise(ys)]
        em = sum(steps) / len(steps) / probe["size"] if len(ys) == 4 else float("nan")
        first = (height - probe["top"] - ys[0]) / probe["size"] if ys else float("nan")
        key = (probe["family"], probe["weight"], probe["size"])
        rows.setdefault(key, {})[probe["kind"]] = (em, first, fonts, len(ys))
        pitch.setdefault(probe["family"], []).append(em)
    for (family, weight, size), kinds in rows.items():
        b_em, b_first, fonts, n = kinds["breaks"]
        p_em, _, _, _ = kinds["paragraphs"]
        note = "" if n == 4 else f"  ({n} baselines!)"
        out.append(
            f"  {family:<16} {weight:<8} {size:>4}  {b_em:>7.4f} {p_em:>7.4f}"
            f"  {b_first:>12.4f}  {', '.join(fonts)}{note}"
        )

    out += ["", f"wrap margin (bold 14 pt label, box = factor x advance sum), {version}:"]
    out.append(f"  {'family':<16} {'label':<28} {'advance pt':>10}  " + "  ".join(FACTORS))
    wraps = {}
    for probe, (ys, _fonts, _h) in zip(probes, pages, strict=True):
        if probe["what"] == "wrap":
            wraps.setdefault((probe["family"], probe["label"], probe["advance"]), {})[
                probe["factor"]
            ] = len(ys)
    for (family, label, advance), counts in wraps.items():
        cells = "  ".join(f"{counts[f]:>5}" for f in FACTORS)
        out.append(f"  {family:<16} {label:<28} {advance:>10.3f}  {cells}")

    out += ["", f"wrap threshold (the smallest box / advance sum that holds one line), {version}:"]
    lines_at = {}
    for (family, label, factor), (ys, _f, _h) in zip(scan_probes, scan_pages, strict=True):
        lines_at.setdefault((family, label), []).append((float(factor), len(ys)))
    for (family, label), seq in lines_at.items():
        one = [f for f, n in seq if n == 1]
        wraps_above = [f for f, n in seq if n > 1 and one and f > min(one)]
        where = f">= {min(one):.4f}" if one else f"> {float(SCAN[-1]):.4f}"
        if one and min(one) == float(SCAN[0]):
            where = f"<= {min(one):.4f} (fits at every scanned factor)"
        flag = "  (not monotonic!)" if wraps_above else ""
        out.append(f"  {family:<16} {label:<28} one line at {where}{flag}")

    # the comparison with §6.4 (24.2): Liberation Sans (Arial) pitch, Arial bold wrap
    arial = pitch["Arial"]
    pitch_ok = all(abs(v - PITCH_EM) <= PITCH_TOLERANCE for v in arial)
    arial_wraps = [c for (fam, _l, _a), c in wraps.items() if fam == "Arial"]
    wrap_ok = all(
        c["1.000"] == 2 and c["1.002"] == 2 and c["1.005"] == 1 and c["1.010"] == 1
        for c in arial_wraps
    )
    spread = {fam: (min(v), max(v)) for fam, v in pitch.items()}
    out += ["", "against spec 002 §6.4 (LibreOffice 24.2):"]
    out.append(
        f"  pitch, Arial -> Liberation Sans: {min(arial):.4f} … {max(arial):.4f} em; "
        f"§6.4 says 1.20 (±{PITCH_TOLERANCE}): {'MATCHES' if pitch_ok else 'DIFFERS'}"
    )
    out.append(
        "  wrap, Arial bold 14 pt: 2 lines at 1.000x and 1.002x, 1 line at 1.005x and 1.010x: "
        f"{'MATCHES' if wrap_ok else 'DIFFERS'}"
    )
    out.append("  pitch per family (min … max em):")
    for fam, (lo, hi) in spread.items():
        out.append(f"    {fam:<16} {lo:.4f} … {hi:.4f}")
    text = "\n".join(out) + "\n"
    print(text, end="")
    if len(argv) > 1:
        Path(argv[1]).write_text(text, encoding="utf-8")
    return 0 if pitch_ok and wrap_ok else 2


if __name__ == "__main__":
    if shutil.which("fc-match") is None:
        sys.exit("stop: fc-match is not installed")
    sys.exit(main(sys.argv))
